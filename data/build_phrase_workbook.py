"""Build the reproducible labelled workbook used by the dashboard demo."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).with_name("hinglish_emotion_phrases.xlsx")


POSITIVE = [
    ("Aaj ka din bahut accha tha 😊", "everyday", "accha"),
    ("Movie mast thi yaar", "entertainment", "mast"),
    ("Khana ekdum tasty tha", "food", "tasty"),
    ("Tumne bahut accha kaam kiya", "work", "accha"),
    ("Finally exam clear ho gaya 🎉", "education", "clear ho gaya"),
    ("Service fast aur helpful thi", "service", "fast; helpful"),
    ("Mujhe ye song bahut pasand hai", "entertainment", "pasand"),
    ("Weekend trip zabardast thi", "travel", "zabardast"),
    ("Gift dekh kar main happy ho gaya", "everyday", "happy"),
    ("Wah kya performance thi 🔥", "entertainment", "wah; performance"),
    ("Team ne great job kiya", "work", "great"),
    ("Coffee perfect bani hai", "food", "perfect"),
    ("Aaj mood ekdum mast hai", "everyday", "mast"),
    ("Thanks yaar tumne help kar di", "everyday", "thanks; help"),
    ("Result expected se better aaya", "education", "better"),
    ("Yeh app use karna easy hai", "technology", "easy"),
    ("Delivery time pe aa gayi 👍", "service", "time pe"),
    ("Mummy ke haath ka khana best hai", "food", "best"),
    ("Interview accha gaya", "work", "accha"),
    ("Itni pyari surprise, love it ❤️", "everyday", "pyari; love"),
    ("Match jeet gaye, maza aa gaya", "sports", "jeet; maza"),
    ("Teacher ne clearly samjhaya", "education", "clearly"),
    ("Phone ka camera amazing hai", "technology", "amazing"),
    ("Bura nahi tha, actually accha laga", "negation", "bura nahi; accha"),
    ("Kaam complete hua, ab sukoon hai", "work", "complete; sukoon"),
]

NEUTRAL = [
    ("Aaj meeting 3 baje hai", "work", "meeting 3 baje"),
    ("Main office pahunch gaya", "work", "office pahunch gaya"),
    ("Package kal deliver hoga", "service", "kal deliver hoga"),
    ("Movie do ghante ki thi", "entertainment", "do ghante"),
    ("Weather thoda cloudy hai", "everyday", "cloudy"),
    ("Train platform 2 par aayegi", "travel", "platform 2"),
    ("Maine email bhej diya", "work", "email bhej diya"),
    ("Class online hogi", "education", "online hogi"),
    ("Phone charging pe laga hai", "technology", "charging"),
    ("Lunch break 1 baje hai", "food", "1 baje"),
    ("Report abhi review mein hai", "work", "review mein"),
    ("Shop shaam 8 baje band hogi", "everyday", "8 baje"),
    ("Aaj traffic normal tha", "travel", "normal"),
    ("Usne message read kiya", "everyday", "message read"),
    ("Ticket confirm ho gaya", "travel", "ticket confirm"),
    ("Meeting reschedule ho gayi", "work", "reschedule"),
    ("Mera order processing mein hai", "service", "processing"),
    ("Kal test ka result aayega", "education", "kal; result"),
    ("Main ghar ja raha hoon", "everyday", "ghar ja raha"),
    ("Theek hai, dekhte hain", "everyday", "theek"),
    ("Price 500 rupees hai", "shopping", "500 rupees"),
    ("Video 10 minute ka hai", "technology", "10 minute"),
    ("Update install ho raha hai", "technology", "install"),
    ("Bus thodi der mein aayegi", "travel", "bus aayegi"),
    ("Aaj Tuesday hai", "everyday", "Tuesday"),
]

NEGATIVE = [
    ("Service bahut kharab thi", "service", "kharab", False),
    ("Movie boring aur bakwaas thi", "entertainment", "boring; bakwaas", False),
    ("Order abhi tak nahi aaya 😡", "service", "nahi aaya", False),
    ("Mujhe ye bilkul pasand nahi hai", "negation", "pasand nahi", False),
    ("Exam result dekh kar sad ho gaya", "education", "sad", False),
    ("Phone baar baar hang ho raha hai", "technology", "hang", False),
    ("Delivery do ghante late thi", "service", "late", False),
    ("Staff ka behaviour rude tha", "service", "rude", False),
    ("Yeh app use karna mushkil hai", "technology", "mushkil", False),
    ("Khana thanda aur bekaar tha", "food", "thanda; bekaar", False),
    ("Mera poora plan cancel ho gaya", "everyday", "cancel", False),
    ("Aaj mood bahut kharab hai", "everyday", "kharab", False),
    ("Worst experience tha", "service", "worst", False),
    ("Network bahut slow chal raha hai", "technology", "slow", False),
    ("Itni ghatiya quality expect nahi thi", "shopping", "ghatiya", False),
    ("Kaam fail ho gaya", "work", "fail", False),
    ("Main is delay se pareshan hoon", "service", "delay; pareshan", False),
    ("Tumne mera trust tod diya", "everyday", "trust tod diya", False),
    ("Ticket book nahi hui", "travel", "nahi hui", False),
    ("Battery bahut jaldi khatam hoti hai", "technology", "jaldi khatam", False),
    ("Meeting bilkul productive nahi thi", "negation", "productive nahi", False),
    ("Wah kya fast service hai, bas do ghante late 😂", "sarcasm", "wah; late", True),
    ("Great, app phir crash ho gayi 😂", "sarcasm", "great; crash", True),
    ("Kya amazing network hai, call har minute drop hoti hai 😂", "sarcasm", "amazing; drop", True),
    ("Bahut accha, mera order phir cancel ho gaya 😂", "sarcasm", "accha; cancel", True),
]


def rows() -> list[dict[str, object]]:
    result: list[dict[str, object]] = []
    for label, phrases, valence in (
        ("positive", POSITIVE, 0.7),
        ("neutral", NEUTRAL, 0.0),
    ):
        for index, (text, category, evidence) in enumerate(phrases, start=1):
            result.append(
                {
                    "text": text,
                    "label": label,
                    "sarcasm": False,
                    "valence": valence,
                    "arousal": 0.45 if any(mark in text for mark in "!😊🎉🔥👍❤️") else 0.2,
                    "evidence": evidence,
                    "pair_id": f"{label[:3]}-{index:02d}",
                    "category": category,
                }
            )
    for index, (text, category, evidence, sarcasm) in enumerate(NEGATIVE, start=1):
        result.append(
            {
                "text": text,
                "label": "negative",
                "sarcasm": sarcasm,
                "valence": -0.8 if sarcasm else -0.7,
                "arousal": 0.75 if sarcasm or any(mark in text for mark in "!😡😂") else 0.4,
                "evidence": evidence,
                "pair_id": f"neg-{index:02d}",
                "category": category,
            }
        )
    return result


def build(output: Path = OUTPUT) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "phrases"
    records = rows()
    headers = list(records[0])
    sheet.append(headers)
    for record in records:
        sheet.append([record[column] for column in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    widths = {"A": 58, "B": 13, "C": 11, "D": 11, "E": 11, "F": 24, "G": 13, "H": 16}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    table = Table(displayName="HinglishEmotionPhrases", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)

    notes = workbook.create_sheet("about")
    notes.append(["Field", "Description"])
    notes_rows = [
        ("Purpose", "Illustrative labelled data for exercising the local project pipeline and dashboard."),
        ("Rows", len(records)),
        ("Balance", "25 positive, 25 neutral, and 25 negative phrases."),
        ("Important", "These hand-labelled examples are a smoke-test set, not a published research benchmark."),
        ("Required columns", "text, label"),
        ("Optional columns", "sarcasm, valence, arousal, evidence, pair_id"),
    ]
    for item in notes_rows:
        notes.append(item)
    for cell in notes[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    notes.column_dimensions[get_column_letter(1)].width = 24
    notes.column_dimensions[get_column_letter(2)].width = 100
    notes.freeze_panes = "A2"

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


if __name__ == "__main__":
    print(build())
