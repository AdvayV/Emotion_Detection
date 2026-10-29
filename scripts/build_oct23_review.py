"""Build the 23 October CodeMood review from the final Markdown and figures.

This script is intentionally separate from the original submission-report
builder. It refuses unresolved evidence markers and requires the supplied
architecture figure before writing the requested DOCX.
"""
from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "oct23review" / "report.md"
FIGURES = ROOT / "docs" / "oct23review" / "figures"
OUTPUT = ROOT / "Documents for review" / "Review reports" / "23OctReview AdvayV KaavishG CD.docx"
PAGE_WIDTH = 8.5
SIDE_MARGIN = 0.75
CONTENT_WIDTH = PAGE_WIDTH - SIDE_MARGIN * 2

FIGURE_RE = re.compile(r"^\*\*Figure\s+(\d+)\s*\.\s*(.*?)\*\*\s*(.*)$", re.I)
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _add_hyperlink(paragraph, label: str, target: str):
    """Add a Word hyperlink for normal external references."""
    relation = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "254E70")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, text: str, *, hyperlinks: bool = True):
    """Render the report's small Markdown inline vocabulary as Word runs."""
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            label, target = LINK_RE.fullmatch(token).groups()
            if hyperlinks and target.startswith(("https://", "http://", "mailto:")):
                _add_hyperlink(paragraph, label, target)
            else:
                paragraph.add_run(label)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _clean_caption(text: str) -> str:
    text = re.sub(r"\s*See\s+\[[^]]*\]\([^)]+\)\.?", "", text, flags=re.I)
    text = IMAGE_RE.sub("", text)
    text = LINK_RE.sub("", text)
    text = text.replace("**", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _figure_source(line: str, following: str = "") -> tuple[str | None, str | None]:
    """Return explicit figure target and alt text from Markdown links/images."""
    text = f"{line} {following}"
    images = IMAGE_RE.findall(text)
    if images:
        return unquote(images[0][1].split("#", 1)[0]), images[0][0]
    for label, target in LINK_RE.findall(text):
        if target.lower().split("#", 1)[0].endswith((".mmd", ".png")):
            return unquote(target.split("#", 1)[0]), label
    return None, None


def _resolve_figure(target: str) -> Path:
    """Resolve figure links without allowing paths outside the review asset dir."""
    raw = Path(target.replace("\\", "/"))
    if raw.suffix.lower() == ".mmd":
        candidate = FIGURES / f"{raw.stem}.png"
    elif raw.suffix.lower() == ".png":
        candidate = (ROOT / raw) if raw.is_absolute() else (SOURCE.parent / raw)
        if candidate.parent.resolve() != FIGURES.resolve():
            candidate = FIGURES / raw.name
    else:
        raise ValueError(f"Figure link must target a .mmd or .png asset: {target}")
    candidate = candidate.resolve()
    if candidate.parent != FIGURES.resolve():
        raise ValueError(f"Figure path is outside the review figures directory: {target}")
    if not candidate.is_file():
        raise FileNotFoundError(f"Required figure image was not rendered: {candidate}")
    with Image.open(candidate) as image:
        image.verify()
    return candidate


def picture(doc: Document, path: Path, *, max_height: float, max_width: float = CONTENT_WIDTH, alt: str = ""):
    with Image.open(path) as image:
        width_px, height_px = image.size
    scale = min(max_width / (width_px / 120), max_height / (height_px / 120))
    width_inches = width_px / 120 * scale
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt or path.stem.replace("_", " "))


def _set_cell_shading(cell, fill: str):
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def _set_cell_borders(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", 90), ("bottom", 90), ("start", 100), ("end", 100)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    props.append(margins)


def _column_widths(rows: list[list[str]], count: int) -> list[float]:
    """Choose readable widths from content rather than assuming a column count."""
    lengths = []
    for col in range(count):
        cell_lengths = [len(row[col]) if col < len(row) else 0 for row in rows]
        longest = max(cell_lengths, default=8)
        lengths.append(max(7, min(50, longest)))
    raw = [value**0.72 for value in lengths]
    minimum = 0.72
    remaining = CONTENT_WIDTH - minimum * count
    if remaining <= 0:
        return [CONTENT_WIDTH / count] * count
    weight = sum(raw) or count
    return [minimum + remaining * value / weight for value in raw]


def table(doc: Document, source_lines: list[str]):
    rows = []
    for line in source_lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if re.fullmatch(r"\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*", line.strip().strip("|")):
            continue
        rows.append(cells)
    if not rows:
        return
    count = max(map(len, rows))
    rows = [row + [""] * (count - len(row)) for row in rows]
    widths = _column_widths(rows, count)
    word_table = doc.add_table(rows=1, cols=count)
    word_table.autofit = False
    word_table.alignment = 1
    for col, width in zip(word_table.columns, widths):
        col.width = Inches(width)
    for ri, row in enumerate(rows):
        cells = word_table.rows[0].cells if ri == 0 else word_table.add_row().cells
        for ci, text in enumerate(row):
            cell = cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_shading(cell, "263F55" if ri == 0 else ("F2F6F8" if ri % 2 else "FFFFFF"))
            _set_cell_borders(cell)
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = Pt(2.5)
            para.paragraph_format.space_after = Pt(2.5)
            para.paragraph_format.line_spacing = 1.0
            inline(para, text)
            for run in para.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if ri == 0 else "202830")
                if ri == 0:
                    run.bold = True
        row_props = word_table.rows[ri]._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_props.append(repeat)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)


def _render_figure(doc: Document, caption_line: str, target: str, alt: str, number: int):
    path = _resolve_figure(target)
    is_architecture = number == 1
    maximum=7.95 if is_architecture else 7.3 if path.stem=="model_comparison" else 6.5 if path.stem=="advanced_fusion" else 5.4 if path.stem=="validation_history" else 4.6
    picture(doc, path, max_height=maximum, alt=alt)
    caption = _clean_caption(caption_line)
    if not caption:
        caption = f"Figure {number}. {alt or path.stem.replace('_', ' ')}."
    para = doc.add_paragraph(style="Caption")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_together = True
    inline(para, caption, hyperlinks=False)


def _style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(PAGE_WIDTH)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(SIDE_MARGIN)
    section.right_margin = Inches(SIDE_MARGIN)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.24)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.7)
    normal.paragraph_format.line_spacing = 1.06
    normal.paragraph_format.space_after = Pt(5.2)
    for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.8)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(11 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    title = doc.styles["Title"]
    title.font.size = Pt(28)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(8)
    subtitle = doc.styles["Subtitle"]
    subtitle.font.size = Pt(14)
    subtitle.font.bold = False
    subtitle.paragraph_format.space_after = Pt(12)
    caption = doc.styles["Caption"]
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.italic = False
    caption.paragraph_format.space_after = Pt(9)
    for style in doc.styles:
        if hasattr(style, "_element"):
            for border in list(style._element.iter(qn("w:pBdr"))):
                border.getparent().remove(border)
    for element in (doc._element, doc.styles._element):
        for border in list(element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)


def _footer(doc: Document):
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("CodeMood  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("66727D")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def _cover(doc: Document):
    p = doc.add_paragraph("Multilingual Emotion Detection", style="Title")
    p.paragraph_format.space_before = Pt(76)
    doc.add_paragraph("Project implementation and training review", style="Subtitle")
    p = doc.add_paragraph("Compiler Design  |  23 October 2026")
    p.paragraph_format.space_before = Pt(14)
    p.runs[0].font.size = Pt(12)
    doc.add_paragraph("Advay Vivek  |  24BCT0136\nKaavish Gogia  |  24BCT0103")
    p = doc.add_paragraph("CodeMood Sarcasm Aware Hinglish Analysis")
    p.paragraph_format.space_before = Pt(10)
    for run in p.runs:
        run.bold = True
    doc.add_paragraph("Review snapshot: 22 October 2026")


def _scan_figures(lines: list[str]) -> dict[int, tuple[int, str, str, str]]:
    figures = {}
    for index, line in enumerate(lines):
        match = FIGURE_RE.match(line.strip())
        if not match:
            continue
        number = int(match.group(1))
        target, alt = _figure_source(line, lines[index + 1] if index + 1 < len(lines) else "")
        if not target:
            raise ValueError(f"Figure {number} needs an explicit .mmd or .png link")
        if number in figures:
            raise ValueError(f"Duplicate Figure {number} caption")
        figures[number] = (index, line.strip(), target, alt or f"Figure {number}")
    return figures


def _render_body(doc: Document, lines: list[str], figures: dict[int, tuple[int, str, str, str]]):
    first_heading = next((i for i, line in enumerate(lines) if line.strip().startswith("# ")), None)
    arch = figures.get(1)
    if arch is None:
        raise ValueError("Report must provide Figure 1 as the full-page architecture diagram")
    # Page 2 is reserved for the architecture figure regardless of source prose order.
    doc.add_page_break()
    doc.add_heading("Project architecture", level=1)
    _render_figure(doc, arch[1], arch[2], arch[3], 1)
    doc.add_page_break()

    i = next((index + 1 for index,line in enumerate(lines) if line.strip()=="<!-- PAGE -->"),0)
    while i < len(lines):
        line = lines[i].strip()
        if not line or line == "<!-- PAGE -->":
            i += 1
            continue
        if first_heading == i and line.startswith("# "):
            i += 1
            continue
        if line.startswith("{{") or "{{FINAL_TESTS}}" in line:
            raise ValueError(f"Unresolved report marker at line {i + 1}: {line}")
        figure_match = FIGURE_RE.match(line)
        if figure_match:
            number = int(figure_match.group(1))
            _, caption, target, alt = figures[number]
            if number != 1:
                _render_figure(doc, caption, target, alt, number)
            # Consume an immediately following Markdown image so it is not duplicated.
            i += 1
            if i < len(lines) and IMAGE_RE.search(lines[i]):
                i += 1
            continue
        if line.startswith("|"):
            group = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                group.append(lines[i].strip())
                i += 1
            table(doc, group)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.startswith("> "):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            inline(para, line[2:])
        elif re.match(r"^[-*+]\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            inline(para, re.sub(r"^[-*+]\s+", "", line))
        elif re.match(r"^\d+[.)]\s+", line):
            para = doc.add_paragraph(style="List Number")
            inline(para, re.sub(r"^\d+[.)]\s+", "", line))
        elif IMAGE_RE.search(line):
            target, alt = _figure_source(line)
            path = _resolve_figure(target)
            picture(doc, path, max_height=4.2, alt=alt or path.stem)
        else:
            para = doc.add_paragraph()
            inline(para, line)
        i += 1


def build():
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Final report source is missing: {SOURCE}")
    source = SOURCE.read_text(encoding="utf-8")
    if "{{FINAL_TESTS}}" in source or re.search(r"\{\{[^}]+\}\}", source):
        raise ValueError("Refusing to build while report placeholders remain unresolved")
    lines = source.splitlines()
    figures = _scan_figures(lines)
    for number, (_, _, target, _) in figures.items():
        _resolve_figure(target)
    doc = Document()
    _style_document(doc)
    doc.core_properties.title = "CodeMood Hinglish Emotion Detection Project Review"
    doc.core_properties.author = "Advay Vivek and Kaavish Gogia"
    doc.core_properties.subject = "Compiler Design project review, 23 October 2026"
    _cover(doc)
    _render_body(doc, lines, figures)
    _footer(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
