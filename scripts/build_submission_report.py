"""Build the submission DOCX from reviewed Markdown and Mermaid figures."""
from pathlib import Path
import re,shutil,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs/submission/report.md'
OUT=ROOT/'Documents for review/Review reports/CodeMood_Hindi_English_Project_Progress_Report.docx'
FIG=ROOT/'docs/submission/figures'


def inline(paragraph,text):
    text=re.sub(r'\[([^\]]+)\]\([^)]*\)',r'\1',text)
    for i,part in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))',text)):
        run=paragraph.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part[1:-1] if part.startswith('*') else part)
        if part.startswith('*') and not part.startswith('**'):run.italic=True
        if part.startswith('**'):run.bold=True
        if part.startswith('`'):run.font.name='Consolas';run.font.size=Pt(9)


def picture(doc,name,max_height=7.9):
    png=FIG/(name+'.png')
    w,h=Image.open(png).size
    width=min(6.95,max_height*w/h)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(4)
    pic=p.add_run().add_picture(str(png),width=Inches(width))
    # Print-quality PNG is embedded; editable SVG and Mermaid sources accompany it.
    pic._inline.docPr.set('descr',name.replace('_',' ')+' architecture diagram; editable Mermaid source supplied')
    p.paragraph_format.keep_with_next=True


def table(doc,lines):
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines if not re.match(r'^\|[\s:|-]+\|$',line)]
    t=doc.add_table(rows=1,cols=len(rows[0]));t.autofit=False
    widths=([1.25,2.45,3.25] if len(rows[0])==3 else [1.6,1.55,2.05,1.75])
    for i,c in enumerate(t.columns):c.width=Inches(widths[i] if i<len(widths) else 6.9/len(rows[0]))
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,text in enumerate(row):
            c=cells[ci];c.width=Inches(widths[ci]);c.vertical_alignment=1
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.05
            inline(p,text)
            for r in p.runs:r.font.size=Pt(9.5);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '202830')
            pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'20384B' if ri==0 else 'F2F6F7' if ri%2 else 'FFFFFF');pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
        trpr=t.rows[ri]._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        if ri==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def build():
    s=SOURCE.read_text(encoding='utf-8')
    if '{{' in s:raise ValueError('Resolve report placeholders before authoring')
    doc=Document()
    sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
    sec.header_distance=Inches(.22);sec.footer_distance=Inches(.22)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        sty=doc.styles[name];sty.font.name='Calibri';sty.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.06;normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Heading 1',16),('Heading 2',12),('Title',30),('Subtitle',15)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=name!='Subtitle'
        doc.styles[name].paragraph_format.space_before=Pt(12);doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=True;doc.styles['Caption'].font.bold=False
    doc.core_properties.title='Multilingual Emotion Detection Research and Implementation Report'
    doc.core_properties.author='Advay Vivek and Kaavish Gogia'
    doc.core_properties.subject='CodeMood project submission'
    p=doc.add_paragraph('Multilingual Emotion Detection',style='Title');p.paragraph_format.space_before=Pt(62)
    doc.add_paragraph('Research and Implementation Report',style='Subtitle')
    p=doc.add_paragraph('CodeMood – Sarcasm-Aware Hinglish Analysis');p.runs[0].bold=True;p.paragraph_format.space_before=Pt(20)
    doc.add_paragraph('Compiler Design\n22 October 2026')
    p=doc.add_paragraph('Advay Vivek\n24BCT0136\n\nKaavish Gogia\n24BCT0103');p.paragraph_format.space_before=Pt(22)
    p=doc.add_paragraph('Local inference · Google Colab training · Seven emotions');p.paragraph_format.space_before=Pt(24)
    doc.add_paragraph('This report explains the implemented pipeline, completed XLM-R and MuRIL experiments, the completed mBERT experiment and the human review workflow. It separates measured results from the work that still requires training or independent evaluation.')
    doc.add_page_break()
    # Full-page system architecture, then connected prose and four module figures.
    doc.add_heading('Complete project architecture',level=1)
    picture(doc,'master_architecture',7.95)
    doc.add_paragraph('Figure 1. Complete CodeMood workflow. The outputs are independent workspace actions. Dashed routes indicate optional review or a future, explicitly approved retraining step. MuRIL is limited to Hindi-English; XLM-R and mBERT support both language mixes.',style='Caption')
    doc.add_page_break()
    body=s.split('<!-- PAGE -->',1)[1]
    lines=body.splitlines();i=0
    figure_names={'Figure 1':'master_architecture','Figure 2':'emotion_inference','Figure 3':'colab_training','Figure 4':'file_evaluation','Figure 5':'disagreement_review'}
    while i<len(lines):
        line=lines[i].strip()
        if not line or line=='<!-- PAGE -->':i+=1;continue
        if line.startswith('|'):
            group=[]
            while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i].strip());i+=1
            table(doc,group);continue
        if line.startswith('**Figure'):
            key=next(k for k in figure_names if line.startswith('**'+k))
            if key!='Figure 1':
                picture(doc,figure_names[key],5.9)
                clean=re.sub(r' See \[.*?\]\(.*?\)\.?','',line).replace('**','')
                doc.add_paragraph(clean,style='Caption')
            i+=1;continue
        if line.startswith('## '):
            text=line[3:];text=re.sub(r'^(\d+)\. ',r'\1  ',text)
            doc.add_heading(text,level=1)
        elif line.startswith('### '):doc.add_heading(line[4:],level=2)
        elif line.startswith('> '):inline(doc.add_paragraph(),line[2:])
        else:
            inline(doc.add_paragraph(),line)
        i+=1
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('CodeMood  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('64748B')
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
    if OUT.exists():
        backup=ROOT/'archive/reports/CodeMood_Report_2026-10-17.docx';backup.parent.mkdir(parents=True,exist_ok=True)
        if not backup.exists():shutil.copyfile(OUT,backup)
    # Remove default template title rules and paragraph borders explicitly.
    for element in [doc._element,doc.styles._element]:
        for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    OUT.parent.mkdir(exist_ok=True);doc.save(OUT)
    print(OUT)

if __name__=='__main__':build()
