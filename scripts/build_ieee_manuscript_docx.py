"""Build an editable Word IEEE-style manuscript from Markdown and references.

The output uses a cover, a two-column body, and a full-width architecture page
at its location in the manuscript. This is IEEE manuscript style, not a
certified IEEE conference template.
"""
from __future__ import annotations

import io
import json
import re
import textwrap
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "docs" / "ieee_manuscript"
SOURCE = BASE / "manuscript.md"
REFERENCES = BASE / "references.json"
FIGURES = BASE / "figures"
OUTPUT = ROOT / "outputs" / "23OctReview AdvayV KaavishG CD IEEE Explained.docx"

LETTER_W, LETTER_H = 8.5, 11.0
MARGIN = 0.75
GUTTER = 0.25
COLUMN_W = (LETTER_W - 2 * MARGIN - GUTTER) / 2
FONT = "Times New Roman"

IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`|\[@[^]]+\]|\[[^\]]+\]\([^)]+\))")
FENCE_RE = re.compile(r"^```(?:text|plaintext|pseudo|pseudocode)?\s*$", re.I)


def _set_font(run, size=10, bold=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_paragraph_font(paragraph, size=10):
    for run in paragraph.runs:
        _set_font(run, size=size)


def _column_settings(section, count: int, gap_inches: float = GUTTER):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(round(gap_inches * 1440)))
    cols.set(qn("w:sep"), "0")


def _format_section(section, *, columns=1, margin=MARGIN):
    section.page_width = Inches(LETTER_W)
    section.page_height = Inches(LETTER_H)
    section.top_margin = Inches(margin)
    section.bottom_margin = Inches(margin)
    section.left_margin = Inches(margin)
    section.right_margin = Inches(margin)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    _column_settings(section, columns)


def _style_document(doc: Document):
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(3.3)
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 10.3), ("Heading 2", 10), ("Heading 3", 9.8)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(7 if name == "Heading 1" else 5)
        style.paragraph_format.space_after = Pt(2.4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    doc.styles["Caption"].font.size = Pt(8.5)
    doc.styles["Caption"].font.italic = False
    doc.styles["Caption"].paragraph_format.space_before = Pt(1)
    doc.styles["Caption"].paragraph_format.space_after = Pt(6)
    # Remove Word template title rules/borders and retain a plain black title.
    for item in (doc._element, doc.styles._element):
        for border in list(item.iter(qn("w:pBdr"))):
            border.getparent().remove(border)


def _set_borders(cell, color="808080", size="4"):
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", 55), ("bottom", 55), ("start", 70), ("end", 70)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    props.append(margins)


def _split_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        content = line.strip().strip("|")
        if re.fullmatch(r"\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*", content):
            continue
        rows.append([cell.strip() for cell in content.split("|")])
    count = max((len(row) for row in rows), default=0)
    return [row + [""] * (count - len(row)) for row in rows]


def _widths(rows: list[list[str]], available: float) -> list[float]:
    count = len(rows[0])
    if not count:
        return []
    measures = []
    for col in range(count):
        longest = max((len(row[col]) for row in rows), default=9)
        measures.append(max(7.0, min(44.0, longest)) ** 0.68)
    floor = min(0.55, available / count * 0.78)
    slack = max(0, available - floor * count)
    total = sum(measures) or count
    return [floor + slack * n / total for n in measures]


TABLE_TITLES = ["Seven emotion checkpoint outputs", "Dataset resources and partitions", "Per class F1 within each model test scope", "Validation subset confidence scaling", "Checking phrase diagnostics", "Matched bilingual synthetic test comparison"]
TABLE_COUNTER = 0

def add_table(doc: Document, lines: list[str], available_width=COLUMN_W):
    global TABLE_COUNTER
    TABLE_COUNTER += 1
    roman = ["I", "II", "III", "IV", "V", "VI"][TABLE_COUNTER-1]
    title = TABLE_TITLES[TABLE_COUNTER-1]
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(3)
    _set_font(caption.add_run("TABLE " + roman + "\n" + title), size=8.3, bold=True)
    rows = _split_table(lines)
    if not rows or not rows[0]:
        return
    widths = _widths(rows, available_width)
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.autofit = False
    t.alignment = 1
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for ri, row in enumerate(rows):
        cells = t.rows[0].cells if ri == 0 else t.add_row().cells
        for ci, text in enumerate(row):
            cell = cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_borders(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = ri < len(rows)-1
            add_inline(p, text)
            for run in p.runs:
                _set_font(run, size=8.2, bold=(ri == 0))
        trpr = t.rows[ri]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            hdr = OxmlElement("w:tblHeader")
            hdr.set(qn("w:val"), "true")
            trpr.append(hdr)
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(1)


def _read_references() -> tuple[dict[str, int], list[dict]]:
    if not REFERENCES.is_file():
        raise FileNotFoundError(f"Reference metadata is missing: {REFERENCES}")
    data = json.loads(REFERENCES.read_text(encoding="utf-8"))
    if isinstance(data, dict) and "references" in data:
        data = data["references"]
    if isinstance(data, dict):
        items = [{"key": key, **(value if isinstance(value, dict) else {"title": str(value)})} for key, value in data.items()]
    elif isinstance(data, list):
        items = data
    else:
        raise ValueError("references.json must be a list or object keyed by citation key")
    keys = {}
    for index, record in enumerate(items, start=1):
        if not isinstance(record, dict):
            raise ValueError("Every reference entry must be an object")
        key = record.get("key") or record.get("id")
        if not key:
            raise ValueError(f"Reference {index} has no key or id")
        if key in keys:
            raise ValueError(f"Duplicate reference key: {key}")
        keys[str(key)] = index
    return keys, items


def add_inline(paragraph, text: str, citations: dict[str, int] | None = None):
    citations = citations or {}
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`"):
            from ieee_equations import INLINE_LATEX, append_inline_symbol
            if token[1:-1] in INLINE_LATEX:
                append_inline_symbol(paragraph, token[1:-1])
                cursor = match.end()
                continue
            run = paragraph.add_run(token[1:-1])
            _set_font(run, size=8.5, name="Consolas")
        elif token.startswith("[@"):
            citation_text = token[2:-1]
            raw_keys = [piece.lstrip("@").strip() for piece in re.split(r"[,;\s]+", citation_text) if piece.lstrip("@").strip()]
            missing = [key for key in raw_keys if key not in citations]
            if missing:
                raise ValueError(f"Citation key(s) not present in references.json: {', '.join(missing)}")
            nums = sorted({citations[key] for key in raw_keys})
            paragraph.add_run("[" + ", ".join(map(str, nums)) + "]")
        else:
            label, _target = LINK_RE.fullmatch(token).groups()
            paragraph.add_run(label)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _image_target(line: str) -> tuple[str, str] | None:
    match = IMAGE_RE.search(line)
    if not match:
        return None
    return match.group(1).strip(), unquote(match.group(2).split("#", 1)[0])


def _resolve_image(target: str) -> Path:
    relative = Path(target.replace("\\", "/"))
    if relative.is_absolute() or relative.suffix.lower() != ".png":
        raise ValueError(f"Figures must be relative PNG files: {target}")
    candidate = (BASE / relative).resolve()
    if candidate.parent != FIGURES.resolve():
        candidate = (FIGURES / relative.name).resolve()
    if candidate.parent != FIGURES.resolve():
        raise ValueError(f"Figure path must stay in {FIGURES}: {target}")
    if not candidate.is_file():
        raise FileNotFoundError(f"Figure PNG is missing: {candidate}")
    with Image.open(candidate) as image:
        image.verify()
    return candidate


def add_grayscale_figure(doc: Document, target: str, caption: str, *, full_width=False):
    path = _resolve_image(target)
    with Image.open(path) as image:
        gray = image.convert("L")
        stream = io.BytesIO()
        gray.save(stream, format="PNG")
        stream.seek(0)
        w, h = gray.size
    max_width = (5.8 if path.stem == "master_architecture" else 6.7) if full_width else COLUMN_W
    max_height = (8.2 if path.stem == "master_architecture" else 4.3) if full_width else 3.0
    scale = min(max_width / (w / 120), max_height / (h / 120))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(stream, width=Inches(w / 120 * scale))
    shape._inline.docPr.set("descr", caption or path.stem.replace("_", " "))
    if caption:
        cp = doc.add_paragraph(style="Caption")
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.keep_together = True
        add_inline(cp, caption)
        for run in cp.runs:
            _set_font(run, size=8.5)


def _pseudocode(doc: Document, text: str):
    """Keep code-like pseudocode editable and compact in the current column."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.06)
    p.paragraph_format.right_indent = Inches(0.04)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    wrapped = []
    for line in text.strip("\n").splitlines():
        expanded = line.expandtabs(4)
        chunks = textwrap.wrap(expanded.strip(), width=50, subsequent_indent="    ",
                               break_long_words=True, break_on_hyphens=False) or [""]
        wrapped.extend(chunks)
    for n, line in enumerate(wrapped):
        if n:
            p.add_run().add_break()
        run = p.add_run(line)
        _set_font(run, size=8, name="Consolas")


def _format_reference(record: dict, number: int) -> str:
    authors = record.get("authors", record.get("author", ""))
    if isinstance(authors, list):
        authors = ", ".join(map(str, authors))
    title = record.get("title", "")
    venue = record.get("venue", record.get("journal", record.get("publisher", "")))
    year = record.get("year", "")
    url = record.get("url", "")
    pieces = [str(authors).strip().rstrip(".") if authors else "", f'"{str(title).strip().rstrip(".")},"' if title else "", str(venue).strip().rstrip(".,") if venue else "", str(year) if year else ""]
    text = ", ".join(piece for piece in pieces if piece)
    if url:
        text += (". " if text and not text.endswith(".") else " ") + str(url)
    elif text and not text.endswith("."):
        text += "."
    return f"[{number}] {text}"


def _cover(doc: Document):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(95)
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("CodeMood: Multilingual Emotion Detection")
    _set_font(run, size=24, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("Prepared by")
    _set_font(run, size=13)
    for text in ("Advay Vivek  |  24BCT0136", "Kaavish Gogia  |  24BCT0103"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        _set_font(p.add_run(text), size=13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    _set_font(p.add_run("Guided by:  Suganthini C"), size=12)
    for text in ("VIT Vellore", "Compiler Design", "23 October 2026"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(p.add_run(text), size=12)


def _insert_page_number(doc: Document):
    for section in [doc.sections[0]]:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        _set_font(run, size=8)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        p._p.append(field)


def _add_display_equation(doc, key, number):
    from ieee_equations import append_equation
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Inches(COLUMN_W / 2), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(COLUMN_W), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    append_equation(p, key, font_size=8.5 if key in {"js", "ece", "weighted_ce", "nll"} else 9.0)
    run=p.add_run("\t(" + str(number) + ")")
    _set_font(run, size=8.5)


def _render_main(doc: Document, lines: list[str], citations: dict[str, int], refs: list[dict]):
    in_code = False
    code_lines: list[str] = []
    equation_number = 0
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if FENCE_RE.match(line):
            if in_code:
                _pseudocode(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(raw.rstrip())
            i += 1
            continue
        if not line:
            i += 1
            continue
        equation = re.fullmatch(r"\[\[equation:([a-z_0-9]+)\]\]", line)
        if equation:
            equation_number += 1
            _add_display_equation(doc, equation.group(1), equation_number)
            i += 1
            continue
        image = _image_target(line)
        if image:
            caption, target = image
            if Path(target).stem == "master_architecture":
                # Return to one column for the architecture at its narrative
                # position, then resume the two-column article on the next page.
                arch_section = doc.add_section(WD_SECTION.NEW_PAGE)
                _format_section(arch_section, columns=1)
                add_grayscale_figure(doc, target, caption, full_width=True)
                main_section = doc.add_section(WD_SECTION.NEW_PAGE)
                _format_section(main_section, columns=2)
                i += 1
                continue
            wide = doc.add_section(WD_SECTION.CONTINUOUS)
            _format_section(wide, columns=1)
            add_grayscale_figure(doc, target, caption, full_width=True)
            resume = doc.add_section(WD_SECTION.CONTINUOUS)
            _format_section(resume, columns=2)
            i += 1
            continue
        if line.startswith("# "):
            # The first Markdown heading is the title already shown on the cover.
            i += 1
            continue
        if line.lower() in {"## references", "## bibliography", "## works cited"}:
            # References are rendered from references.json in the supplied order.
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("## "):
                i += 1
            continue
        if line.startswith("|"):
            group = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                group.append(lines[i].strip())
                i += 1
            add_table(doc, group)
            continue
        if line.startswith("## "):
            heading = doc.add_heading(line[3:].strip(), level=1)
            for run in heading.runs: _set_font(run, size=10.3 if 1 == 1 else 10, bold=True)
        elif line.startswith("### "):
            heading = doc.add_heading(line[4:].strip(), level=2)
            for run in heading.runs: _set_font(run, size=10.3 if 2 == 1 else 10, bold=True)
        elif line.startswith("#### "):
            heading = doc.add_heading(line[5:].strip(), level=3)
            for run in heading.runs: _set_font(run, size=10.3 if 3 == 1 else 10, bold=True)
        elif re.match(r"^[-*+]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*+]\s+", "", line), citations)
            _set_paragraph_font(p, 10)
        elif re.match(r"^\d+[.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+[.)]\s+", "", line), citations)
            _set_paragraph_font(p, 10)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.16)
            add_inline(p, line[2:], citations)
            _set_paragraph_font(p, 9.5)
        else:
            p = doc.add_paragraph()
            add_inline(p, line, citations)
            _set_paragraph_font(p, 10)
            if line.startswith(("**Algorithm", "**Worked example")):
                p.paragraph_format.keep_with_next = True
        i += 1
    if in_code:
        raise ValueError("Unclosed Markdown pseudocode fence")
    heading = doc.add_heading("References", level=1)
    for run in heading.runs: _set_font(run, size=10.3, bold=True)
    for n, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.keep_together = True
        run = p.add_run(_format_reference(ref, n))
        _set_font(run, size=8.5)


def build():
    if not SOURCE.is_file():
        raise FileNotFoundError(f"Manuscript source is missing: {SOURCE}")
    if not REFERENCES.is_file():
        raise FileNotFoundError(f"Reference JSON is missing: {REFERENCES}")
    markdown = SOURCE.read_text(encoding="utf-8")
    if re.search(r"\{\{[^}]+\}\}", markdown):
        raise ValueError("Refusing to build while unresolved manuscript placeholders remain")
    citations, refs = _read_references()
    lines = markdown.splitlines()
    architecture_line = next((line for line in lines if (image := _image_target(line)) and Path(image[1]).stem == "master_architecture"), None)
    if architecture_line is None:
        raise ValueError("manuscript.md must include ![caption](figures/master_architecture.png)")
    _resolve_image(_image_target(architecture_line)[1])

    doc = Document()
    _style_document(doc)
    first = doc.sections[0]
    _format_section(first, columns=1)
    _cover(doc)

    main = doc.add_section(WD_SECTION.NEW_PAGE)
    _format_section(main, columns=2)
    _render_main(doc, lines, citations, refs)
    doc.core_properties.title = "CodeMood Multilingual Emotion Detection"
    doc.core_properties.author = "Advay Vivek; Kaavish Gogia"
    doc.core_properties.subject = "Compiler Design manuscript, 23 October 2026"
    _insert_page_number(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
